"""Foodspring 모듈 단위 테스트 (네트워크 미사용)."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from app.core.base import Control
from app.core.foodspring.engine import FoodSpringCrawler
from app.core.foodspring.exporter import FoodSpringExporter
from app.core.foodspring.outcome import RunOutcome, determine_outcome
from app.core.foodspring.preflight import (
    PreflightStatus,
    check_runtime,
)
from app.models.foodspring_records import (
    FoodSpringRunConfig,
    FoodSpringRunSummary,
)


class FoodSpringConfigTest(unittest.TestCase):
    def test_defaults_valid(self):
        cfg = FoodSpringRunConfig(output_dir=Path("/tmp/x"))
        self.assertEqual(cfg.max_pages, 200)
        self.assertEqual(cfg.workers, 3)

    def test_invalid_workers(self):
        with self.assertRaises(ValueError):
            FoodSpringRunConfig(output_dir=Path("/tmp/x"), workers=0)

    def test_invalid_delay(self):
        with self.assertRaises(ValueError):
            FoodSpringRunConfig(output_dir=Path("/tmp/x"), delay_min=2.0, delay_max=1.0)

    def test_windows_forbidden_prefix(self):
        with self.assertRaises(ValueError):
            FoodSpringRunConfig(output_dir=Path("/tmp/x"), output_prefix="a:b")

    def test_limit_products(self):
        with self.assertRaises(ValueError):
            FoodSpringRunConfig(output_dir=Path("/tmp/x"), limit_products=0)


class FoodSpringOutcomeTest(unittest.TestCase):
    def test_success(self):
        s = FoodSpringRunSummary(xlsx_path="/tmp/a.xlsx", termination_reason="complete")
        self.assertEqual(determine_outcome(s), RunOutcome.SUCCESS)

    def test_save_error_priority(self):
        s = FoodSpringRunSummary(save_error="disk full", xlsx_path="/tmp/a.xlsx")
        self.assertEqual(determine_outcome(s), RunOutcome.SAVE_ERROR)

    def test_cancelled(self):
        s = FoodSpringRunSummary(cancelled=True)
        self.assertEqual(determine_outcome(s), RunOutcome.CANCELLED)

    def test_error(self):
        s = FoodSpringRunSummary(error="boom", termination_reason="error")
        self.assertEqual(determine_outcome(s), RunOutcome.ERROR)

    def test_no_records(self):
        s = FoodSpringRunSummary(termination_reason="no_items")
        self.assertEqual(determine_outcome(s), RunOutcome.NO_RECORDS)


class FoodSpringPreflightTest(unittest.TestCase):
    def test_runtime_ok(self):
        result = check_runtime()
        self.assertEqual(result.status, PreflightStatus.OK)


class FoodSpringExporterTest(unittest.TestCase):
    def test_excel_sheets(self):
        from openpyxl import load_workbook

        products = {
            "1001": {
                "nid": 1001,
                "name": "테스트 상품",
                "price": {"salePrice": 1000, "discountRate": 10},
                "images": {"primaryUrl": "http://img/1.png"},
                "vendor": {"nid": 77, "name": "테스트셀러"},
            }
        }
        infos = {
            "77": {
                "seller_id": 77,
                "store_name": "테스트셀러",
                "owner_name": "김대표",
                "business_number": "123-45-67890",
                "phone": "010-1111-2222",
                "email": "seller@example.com",
                "address": "서울시 테스트구",
                "ecommerce_report_number": "2020-테스트-1",
                "customer_service_number": "02-1111-2222",
                "_product_count": 1,
                "_representative_pid": "1001",
            }
        }
        with tempfile.TemporaryDirectory() as tmp:
            cfg = FoodSpringRunConfig(output_dir=Path(tmp), output_prefix="test_out")
            path = FoodSpringExporter().save(products, infos, cfg)
            self.assertTrue(path.exists())
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["셀러정보", "상품목록"])
            ws = wb["셀러정보"]
            self.assertEqual(ws.max_row, 2)  # 헤더 + 1
            self.assertEqual(ws.cell(row=2, column=4).value, "123-45-67890")
            self.assertEqual(ws.cell(row=2, column=3).value, "김대표")
            self.assertEqual(ws.cell(row=2, column=6).value, "seller@example.com")
            ws2 = wb["상품목록"]
            self.assertEqual(ws2.max_row, 2)
            self.assertEqual(ws2.cell(row=2, column=2).value, "테스트 상품")
            self.assertEqual(ws2.cell(row=2, column=11).value, "seller@example.com")


class FoodSpringEngineTest(unittest.TestCase):
    """오프라인 엔진 테스트: fetcher/exporter 를 모의(mock) 주입."""

    def _fake_fetcher(self, url):
        class FakePage:
            cookies = (
                {"name": "FS_TOKEN", "value": "fake_token"},
                {"name": "areaId", "value": "913"},
            )

            def __init__(self) -> None:
                self.status = 200

        return FakePage()

    def _mock_session(self, crawler, edge_nodes):
        """GraphQL mock: 목록(goodsList) + 셀러(node) 쿼리 모두 처리."""
        class FakeResp:
            def __init__(self, payload) -> None:
                self._payload = payload

            def raise_for_status(self):
                return None

            def json(self):
                return self._payload

        calls = {"n": 0}

        def _cursor(idx: int) -> str:
            import base64

            return base64.b64encode(f"custom-cursor{idx}".encode()).decode()

        def _vendor_node(node_id: str) -> dict:
            sid = node_id.replace("vendor_", "")
            return {
                "data": {
                    "node": {
                        "__typename": "Vendor",
                        "nid": int(sid),
                        "name": f"셀러{sid}",
                        "ownerName": f"대표{sid}",
                        "businessRegistrationNumber": f"111-11-{sid}0",
                        "mailOrderRegistrationNumber": "신고번호",
                        "businessAddress": "주소",
                        "contact": "010-0000-0000",
                        "email": f"seller{sid}@example.com",
                        "customerServiceNumber": "02-0000-0000",
                        "id": node_id,
                    }
                }
            }

        def fake_post(url, **kwargs):
            import base64

            body = json.loads(kwargs["data"].decode("utf-8"))
            query = body.get("query", "")
            if "node(id" in query:
                return FakeResp(_vendor_node(body["variables"]["id"]))
            after = body["variables"]["after"]
            page_size = len(edge_nodes[0])
            if after is None:
                idx = 0
            else:
                raw = base64.b64decode(after).decode()
                idx = (int(raw.split("custom-cursor")[1]) + 1) // page_size
            if idx >= len(edge_nodes):
                gl = {"pageInfo": {"hasNextPage": False}, "edges": []}
            else:
                last = idx * page_size + len(edge_nodes[idx]) - 1
                gl = {
                    "pageInfo": {
                        "hasNextPage": idx + 1 < len(edge_nodes),
                        "endCursor": _cursor(last),
                    },
                    "edges": [{"node": n} for n in edge_nodes[idx]],
                }
            calls["n"] += 1
            return FakeResp({"data": {"goodsList": gl}})

        crawler._session.post = fake_post

    def test_extract_vendor_from_next_data(self):
        html = (
            '<html><script id="__NEXT_DATA__" type="application/json">'
            + json.dumps({
                "props": {
                    "pageProps": {
                        "initialRecords": {
                            "vendor_3575": {
                                "__typename": "Vendor",
                                "nid": 3575,
                                "name": "BEST RICE 31",
                                "businessRegistrationNumber": "206-93-80114",
                                "contact": "010-5042-9982",
                                "businessAddress": "서울 서초구 양재동 223",
                                "mailOrderRegistrationNumber": "2020-서울서초-4351",
                                "customerServiceNumber": "02-3461-3134",
                            }
                        }
                    }
                }
            })
            + "</script></html>"
        )
        info = FoodSpringCrawler._extract_vendor(html, "3575")
        self.assertIsNotNone(info)
        self.assertEqual(info["business_number"], "206-93-80114")
        self.assertEqual(info["store_name"], "BEST RICE 31")

    def test_full_run_mocked(self):
        edges = [
            [{
                "nid": 1000 + i,
                "name": f"상품{i}",
                "price": {"salePrice": 5000, "discountRate": 10},
                "images": {"primaryUrl": "http://img/x.png"},
                "vendor": {"nid": 10 + i, "name": f"셀러{i}"},
            } for i in range(5)],
            [{
                "nid": 2000 + i,
                "name": f"상품B{i}",
                "price": {"salePrice": 3000, "discountRate": 20},
                "images": {"primaryUrl": "http://img/y.png"},
                "vendor": {"nid": 20 + i, "name": f"셀러B{i}"},
            } for i in range(5)],
        ]

        with tempfile.TemporaryDirectory() as tmp:
            cfg = FoodSpringRunConfig(
                output_dir=Path(tmp),
                output_prefix="mocked",
                max_pages=3,
                delay_min=0.0,
                delay_max=0.0,
            )
            crawler = FoodSpringCrawler(
                cfg,
                Control(),
                fetcher=self._fake_fetcher,
                exporter=FoodSpringExporter(),
            )
            self._mock_session(crawler, edges)

            # 상세 HTML mock: 셀러별 데이터 반환 (pid 기반)
            def fake_get(url, **kwargs):
                import re as _re

                m = _re.search(r"/goods/detail/(\d+)", url)
                pid = int(m.group(1))
                sid = None
                for e in edges:
                    for n in e:
                        if n["nid"] == pid:
                            sid = n["vendor"]["nid"]
                html = (
                    '<html><script id="__NEXT_DATA__" type="application/json">'
                    + json.dumps({
                        "props": {
                            "pageProps": {
                                "initialRecords": {
                                    f"vendor_{sid}": {
                                        "__typename": "Vendor",
                                        "nid": sid,
                                        "name": f"셀러{sid}",
                                        "businessRegistrationNumber": f"111-11-{sid:05d}",
                                        "contact": "010-0000-0000",
                                        "businessAddress": "주소",
                                        "mailOrderRegistrationNumber": "신고번호",
                                        "customerServiceNumber": "02-0000-0000",
                                    }
                                }
                            }
                        }
                    })
                    + "</script></html>"
                )

                class FakeResp:
                    def raise_for_status(self):
                        return None

                    def __init__(self, html) -> None:
                        self.encoding = "utf-8"
                        self.text = html

                return FakeResp(html)

            crawler._session.get = fake_get
            crawler._worker_session = lambda: crawler._session
            crawler._session.headers["Cookie"] = "FS_TOKEN=fake"

            logs = []
            crawler.on_log = lambda msg: logs.append(msg)
            summary = crawler.run()

            self.assertFalse(summary.cancelled, summary.error)
            self.assertIsNone(summary.error)
            self.assertEqual(summary.products_seen, 10)
            self.assertEqual(summary.unique_vendors, 10)
            self.assertEqual(summary.business_info_success, 10)
            self.assertEqual(summary.email_success, 10)
            self.assertEqual(summary.termination_reason, "complete")
            self.assertIsNotNone(summary.xlsx_path)
            self.assertTrue(Path(summary.xlsx_path).exists())
            record = summary.records[0]
            self.assertTrue(record["email"].endswith("@example.com"))
            self.assertTrue(record["owner_name"].startswith("대표"))

    def test_cancel_saves_partial(self):
        """셀러 수집 중 취소 시 지금까지 수집된 분량이 부분 엑셀로 저장된다."""
        edges = [[
            {
                "nid": 1000 + i,
                "name": f"상품{i}",
                "price": {"salePrice": 5000, "discountRate": 10},
                "images": {"primaryUrl": "http://img/x.png"},
                "vendor": {"nid": 100 + i, "name": f"셀러{i}"},
            } for i in range(12)
        ]]
        with tempfile.TemporaryDirectory() as tmp:
            cfg = FoodSpringRunConfig(
                output_dir=Path(tmp),
                output_prefix="cancel_test",
                max_pages=3,
                workers=1,
                delay_min=0.0,
                delay_max=0.0,
            )
            crawler = FoodSpringCrawler(
                cfg, Control(), fetcher=self._fake_fetcher, exporter=FoodSpringExporter()
            )
            self._mock_session(crawler, edges)
            crawler._session.get = self._fake_get_for_sellers(edges)
            crawler._worker_session = lambda: crawler._session
            crawler._session.headers["Cookie"] = "FS_TOKEN=fake"

            # 셀러 3건 수집 시점에 취소 요청
            def cancel_hook(kind, current, total):
                if kind == "seller" and current >= 3:
                    crawler.control.request_cancel()

            crawler.on_progress = cancel_hook
            summary = crawler.run()

            self.assertTrue(summary.cancelled)
            self.assertIsNone(summary.error)
            self.assertEqual(summary.unique_products, 12)
            self.assertGreaterEqual(summary.unique_vendors, 2)
            self.assertLess(summary.unique_vendors, 12)
            # 부분 엑셀이 저장되어야 함
            self.assertIsNotNone(summary.xlsx_path)
            self.assertTrue(Path(summary.xlsx_path).exists())

    def test_save_error_mapped(self):
        """엑셀 저장 실패가 save_error + SAVE_ERROR 로 매핑된다."""
        edges = [[{
            "nid": 1, "name": "a", "price": {"salePrice": 1},
            "images": {}, "vendor": {"nid": 5, "name": "v"},
        }]]

        class BrokenExporter:
            def save(self, products, infos, config):
                raise OSError("disk full")

        with tempfile.TemporaryDirectory() as tmp:
            cfg = FoodSpringRunConfig(
                output_dir=Path(tmp), output_prefix="save_err",
                max_pages=3, delay_min=0.0, delay_max=0.0,
            )
            crawler = FoodSpringCrawler(
                cfg, Control(), fetcher=self._fake_fetcher, exporter=BrokenExporter()
            )
            self._mock_session(crawler, edges)
            crawler._session.get = self._fake_get_for_sellers(edges)
            crawler._worker_session = lambda: crawler._session
            crawler._session.headers["Cookie"] = "FS_TOKEN=fake"

            summary = crawler.run()
            self.assertEqual(determine_outcome(summary), RunOutcome.SAVE_ERROR)
            self.assertIsNotNone(summary.save_error)

    def test_list_consecutive_errors_fails(self):
        """목록 페이지 연속 실패가 정상 완료가 아닌 실패로 판정된다."""
        with tempfile.TemporaryDirectory() as tmp:
            cfg = FoodSpringRunConfig(
                output_dir=Path(tmp), output_prefix="list_err",
                max_pages=20, delay_min=0.0, delay_max=0.0,
            )
            crawler = FoodSpringCrawler(
                cfg, Control(), fetcher=self._fake_fetcher, exporter=FoodSpringExporter()
            )

            class Boom(Exception):
                pass

            def fake_post(url, **kwargs):
                raise Boom("network down")

            crawler._session.post = fake_post
            crawler._session.headers["Cookie"] = "FS_TOKEN=fake"
            crawler._retry_delay = 0.0

            summary = crawler.run()
            self.assertIsNotNone(summary.error)
            self.assertNotEqual(summary.termination_reason, "complete")
            self.assertEqual(summary.products_seen, 0)

    def _fake_get_for_sellers(self, edges):
        """모든 상세 요청에 vendor 정보가 있는 HTML을 반환하는 mock get."""
        def fake_get(url, **kwargs):
            import re as _re

            m = _re.search(r"/goods/detail/(\d+)", url)
            pid = int(m.group(1))
            sid = None
            for e in edges:
                for n in e:
                    if n["nid"] == pid:
                        sid = n["vendor"]["nid"]
            html = (
                '<html><script id="__NEXT_DATA__" type="application/json">'
                + json.dumps({
                    "props": {"pageProps": {"initialRecords": {
                        f"vendor_{sid}": {
                            "__typename": "Vendor", "nid": sid, "name": f"셀러{sid}",
                            "businessRegistrationNumber": f"111-11-{sid:05d}",
                            "contact": "010-0000-0000", "businessAddress": "주소",
                            "mailOrderRegistrationNumber": "신고번호",
                            "customerServiceNumber": "02-0000-0000",
                        }
                    }}}})
                + "</script></html>"
            )

            class FakeResp:
                def raise_for_status(self):
                    return None

                def __init__(self, html) -> None:
                    self.encoding = "utf-8"
                    self.text = html

            return FakeResp(html)

        return fake_get


if __name__ == "__main__":
    unittest.main()