"""Foodspring 결과 저장: Excel(xlsx) 원자적 쓰기."""

from __future__ import annotations

import os
import sys
import tempfile
import threading
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

from app.models.foodspring_records import FoodSpringRunConfig

_EXPORT_LOCK_FILENAME = ".foodspring_export.lock"
_save_guard = threading.Lock()

# 상품 시트 컬럼: (내부키, 헤더)
PRODUCT_SHEET_COLS = [
    ("url", "상품 URL"),
    ("name", "상품명"),
    ("sale_price", "판매가(원)"),
    ("discount_rate", "할인율(%)"),
    ("image_url", "이미지 URL"),
    ("seller_id", "셀러ID"),
    ("seller_name", "셀러명(스토어)"),
    ("owner_name", "대표자명"),
    ("business_number", "사업자등록번호"),
    ("phone", "연락처"),
    ("email", "이메일"),
    ("address", "사업장 소재지"),
    ("ecommerce_report_number", "통신판매 신고번호"),
    ("customer_service_number", "고객센터 전화"),
]

# 셀러 시트 컬럼
SELLER_SHEET_COLS = [
    ("seller_id", "셀러ID"),
    ("store_name", "셀러명(스토어)"),
    ("owner_name", "대표자명"),
    ("business_number", "사업자등록번호"),
    ("phone", "연락처"),
    ("email", "이메일"),
    ("address", "사업장 소재지"),
    ("ecommerce_report_number", "통신판매 신고번호"),
    ("customer_service_number", "고객센터 전화"),
    ("product_count", "판매 상품수"),
    ("url", "대표 상품 URL"),
]


@contextmanager
def _exclusive_save_lock(out_dir: Path):
    lock_path = out_dir / _EXPORT_LOCK_FILENAME
    try:
        lock_file = open(lock_path, "a+b")  # noqa: SIM115
    except OSError as e:
        raise OSError(f"Foodspring 결과 잠금 파일을 만들 수 없습니다: {lock_path}: {e}") from e
    try:
        lock_file.seek(0, os.SEEK_END)
        if lock_file.tell() == 0:
            lock_file.write(b"\0")
            lock_file.flush()
        lock_file.seek(0)
        if sys.platform == "win32":
            import msvcrt

            msvcrt.locking(lock_file.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl

            fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError as e:
        lock_file.close()
        raise OSError(
            f"다른 프로세스가 '{out_dir}'에 Foodspring 결과를 저장 중입니다. "
            "저장이 끝난 뒤 다시 시도하세요."
        ) from e
    try:
        yield
    finally:
        lock_file.close()


def _seller_sort_key(item):
    """숫자 셀러ID 우선, 비숫자는 뒤로 정렬 (ValueError 방어)."""
    try:
        return (0, int(item[0]))
    except (ValueError, TypeError):
        return (1, str(item[0]))


class FoodSpringExporter:
    def save(self, products: dict, infos: dict, config: FoodSpringRunConfig) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        out_dir = config.output_dir
        out_dir.mkdir(parents=True, exist_ok=True)
        # 의도적으로 로컬 시각 파일명 유지 — 기존 출력 파일명 컨벤션과 동일
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")  # noqa: DTZ005
        prefix = config.output_prefix or f"foodspring_wcpd_{ts}"
        xlsx_path = out_dir / f"{prefix}.xlsx"

        with _save_guard, _exclusive_save_lock(out_dir):
            tmp_path = None
            try:
                fd, tmp_path = tempfile.mkstemp(
                    suffix=".xlsx", prefix=".foodspring_", dir=str(out_dir)
                )
                os.close(fd)

                wb = Workbook()
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill("solid", fgColor="2F5496")

                # ── 셀러 시트 ──
                ws_seller = wb.active
                ws_seller.title = "셀러정보"
                for c, (_, label) in enumerate(SELLER_SHEET_COLS, 1):
                    cell = ws_seller.cell(row=1, column=c, value=label)
                    cell.font = header_font
                    cell.fill = header_fill
                for r, (sid, info) in enumerate(
                    sorted(infos.items(), key=_seller_sort_key), 2
                ):
                    info = info or {}
                    pid = info.get("_representative_pid", "")
                    row = [
                        info.get("seller_id") or sid,
                        info.get("store_name", ""),
                        info.get("owner_name", ""),
                        info.get("business_number", ""),
                        info.get("phone", ""),
                        info.get("email", ""),
                        info.get("address", ""),
                        info.get("ecommerce_report_number", ""),
                        info.get("customer_service_number", ""),
                        info.get("_product_count", ""),
                        (
                            f"https://www.foodspring.co.kr/goods/detail/{pid}"
                            if pid else ""
                        ),
                    ]
                    for c, val in enumerate(row, 1):
                        ws_seller.cell(row=r, column=c, value=val)
                for c in range(1, len(SELLER_SHEET_COLS) + 1):
                    ws_seller.column_dimensions[get_column_letter(c)].width = 22

                # ── 상품 시트 ──
                ws_prod = wb.create_sheet("상품목록")
                for c, (_, label) in enumerate(PRODUCT_SHEET_COLS, 1):
                    cell = ws_prod.cell(row=1, column=c, value=label)
                    cell.font = header_font
                    cell.fill = header_fill

                r = 2
                for nid in sorted(products.keys(), key=lambda x: int(x)):
                    node = products[nid]
                    sid = (node.get("vendor") or {}).get("nid")
                    info = infos.get(str(sid)) or {}
                    price = node.get("price") or {}
                    img = (node.get("images") or {}).get("primaryUrl") or ""
                    row = [
                        f"https://www.foodspring.co.kr/goods/detail/{nid}",
                        node.get("name", ""),
                        price.get("salePrice"),
                        price.get("discountRate"),
                        img,
                        sid,
                        (node.get("vendor") or {}).get("name", ""),
                        info.get("owner_name", "") if info else "",
                        info.get("business_number", "") if info else "",
                        info.get("phone", "") if info else "",
                        info.get("email", "") if info else "",
                        info.get("address", "") if info else "",
                        info.get("ecommerce_report_number", "") if info else "",
                        info.get("customer_service_number", "") if info else "",
                    ]
                    for c, val in enumerate(row, 1):
                        ws_prod.cell(row=r, column=c, value=val)
                    r += 1
                for c in range(1, len(PRODUCT_SHEET_COLS) + 1):
                    width = 60 if c == 2 else (45 if c in (1, 5) else 14)
                    ws_prod.column_dimensions[get_column_letter(c)].width = width
                ws_prod.freeze_panes = "A2"
                ws_seller.freeze_panes = "A2"

                wb.save(tmp_path)
                os.replace(tmp_path, xlsx_path)
                tmp_path = None
            finally:
                if tmp_path is not None and os.path.exists(tmp_path):
                    try:
                        os.unlink(tmp_path)
                    except OSError:
                        pass

        return xlsx_path